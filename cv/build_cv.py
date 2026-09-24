#!/usr/bin/env python3
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.style import WD_STYLE_TYPE

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)
DARK = RGBColor(0x11, 0x11, 0x11)
TXT = RGBColor(0x22, 0x22, 0x22)
LINK = RGBColor(0x11, 0x55, 0xCC)
FONT = "Calibri"

doc = Document()

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.0)
sec.bottom_margin = Cm(1.0)
sec.left_margin = Cm(1.2)
sec.right_margin = Cm(1.2)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)

def set_border(paragraph, color="1F3864", sz=4):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_hyperlink(paragraph, url, text, color="1155CC", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), FONT); rFonts.set(qn('w:hAnsi'), FONT)
    rPr.append(rFonts)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '18'); rPr.append(sz)
    c = OxmlElement('w:color'); c.set(qn('w:val'), color); rPr.append(c)
    if underline:
        u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement('w:t'); t.text = text; new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def section_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(5)
    set_border(p)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.color.rgb = NAVY
    r.font.size = Pt(10.5)
    return p

def role_line(role, dates):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(18.5), WD_TAB_ALIGNMENT.RIGHT)
    r1 = p.add_run(role)
    r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = DARK
    r2 = p.add_run("\t" + dates)
    r2.italic = True; r2.font.size = Pt(9.5); r2.font.color.rgb = GREY
    return p

def org_line(org):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(org)
    r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = NAVY
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.55)
    for run in p.runs:
        run.font.size = Pt(9.5)
    r = p.add_run(text)
    r.font.size = Pt(9.5); r.font.color.rgb = TXT; r.font.name = FONT
    return p

def plain(text, after=3, size=9.5, italic=False, bold=False, color=TXT):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.color.rgb = color; r.italic = italic; r.bold = bold
    return p

def edu_line(degree, dates):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(18.5), WD_TAB_ALIGNMENT.RIGHT)
    r1 = p.add_run(degree)
    r1.bold = True; r1.font.size = Pt(9.5); r1.font.color.rgb = DARK
    r2 = p.add_run("\t" + dates)
    r2.italic = True; r2.font.size = Pt(9); r2.font.color.rgb = GREY
    return p

def edu_org(org):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4.5)
    r = p.add_run(org)
    r.font.size = Pt(9); r.font.color.rgb = GREY
    return p

def skill_line(cat, items):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.5)
    r1 = p.add_run(cat + " : ")
    r1.bold = True; r1.font.size = Pt(9.5); r1.font.color.rgb = DARK
    r2 = p.add_run(items)
    r2.font.size = Pt(9.5); r2.font.color.rgb = TXT
    return p

# Header
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("CHEIKH DAROU BEYE")
r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
r = p.add_run("Data Scientist & Développeur Full-Stack")
r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

p = doc.add_paragraph()
set_border(p, color="1F3864", sz=8)
p.paragraph_format.space_after = Pt(7)
r = p.add_run("Avignon, France  |  +33 6 61 74 78 96  |  cheikhdaroubeye@gmail.com")
r.font.size = Pt(9); r.font.color.rgb = GREY

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
add_hyperlink(p, "https://sen-donnees.github.io/portfolio-cdb/", "Portfolio")
r = p.add_run("   |   ")
r.font.size = Pt(9); r.font.color.rgb = GREY
add_hyperlink(p, "https://www.linkedin.com/in/cheikh-darou-beye-9544991b8/", "LinkedIn")
r = p.add_run("   |   ")
r.font.size = Pt(9); r.font.color.rgb = GREY
add_hyperlink(p, "https://github.com/sen-donnees", "GitHub")

section_heading("Profil")
plain("Data Scientist et Développeur Full-Stack. Je conçois des produits de bout en bout — de l'analyse statistique et du machine learning jusqu'à des plateformes web en production. Fondateur ou cofondateur de plusieurs produits actifs (recrutement, génération de documents, sécurité privée), formation solide en statistiques (Master 2, Université Grenoble Alpes) et expérience concrète en gestion de projet data.", after=4)

section_heading("Expérience professionnelle")

role_line("Cofondateur & Directeur Technique (CTO)", "2026 – présent")
org_line("YoonuJob — Plateforme de recrutement, Avignon")
bullet("Conception et développement de l'intégralité de la plateforme : candidats, entreprises, candidatures, générateur de CV (3 modèles, sans compte requis).")
bullet("Architecture Django 5.2 / PostgreSQL / Tailwind CSS, déploiement Docker multi-étapes.")
bullet("Back-office sur mesure (~19 écrans) pour la modération et la gestion de la plateforme.")
bullet("300+ tests automatisés ; sécurité applicative (limitation de débit, protection anti-bruteforce, nettoyage des entrées utilisateur).")
bullet("Intégration de DocuSen (voir ci-dessous) dans la plateforme.")

role_line("Fondateur & Développeur unique", "2025 – 2026 (intégré à YoonuJob)")
org_line("DocuSen — Génération de documents professionnels")
bullet("Produit web de génération de documents (factures, contrats, fiches de paie, attestations) : 27 types de documents dans 5 familles, export PDF.")
bullet("693 tests automatisés, intégration continue (GitHub Actions), déploiement Docker sur Render.")
bullet("Stack : Python, Django, WeasyPrint, SQLite, Bootstrap 5.")

role_line("Cofondateur & Développeur", "2026 – présent")
org_line("Sentazur — Société de sécurité privée, Avignon")
bullet("Développement du site public et du back-office complet (13 modules planifiés : clients, devis, agents...) en tant que responsable technique.")
bullet("102 tests automatisés ; système d'alertes automatiques avant expiration des qualifications des agents (ex. certification SSIAP).")
bullet("Stack : Django, PostgreSQL, Tailwind CSS, WeasyPrint.")

role_line("Fondateur & Développeur unique", "2026")
org_line("SenBioStat — Application web d'analyse statistique")
bullet("Outil en ligne exécutant des tests statistiques (Khi², t-test, ANOVA, régression, Kaplan-Meier) avec interprétation automatique en français et export PDF (WeasyPrint).")

role_line("Développeur freelance", "2026")
org_line("Sites clients — AJEL de Rufisque (club de football) & Commune de Keur Moussa")
bullet("Développement de bout en bout de deux sites Django pour des clients externes : club de football professionnel (actualités, effectif, boutique, sponsors) et collectivité locale (actualités, projets, annuaire des villages).")

role_line("Chef de Projet Data Analyst", "Mars 2023 – Mai 2024")
org_line("ITSAP – Institut de l'Abeille (INRAE), Avignon")
bullet("Plateforme de données pour balances de ruches connectées : nettoyage et validation des données terrain.")
bullet("Développement d'un algorithme de régression segmentée (R) pour détecter automatiquement les périodes de miellée.")
bullet("Application interactive R Shiny pour la visualisation en temps réel ; indicateurs géolocalisés et reporting.")

role_line("Data Analyst (Stage)", "Mai 2022 – Août 2022")
org_line("Marie Thieulin, Grenoble")
bullet("Analyse de données énergétiques mondiales (production et consommation) et création de tableaux de bord interactifs avec Power BI et Qlik Sense.")

section_heading("Projets académiques")

role_line("Détection de nouveaux sous-types microbiens par apprentissage de métrique", "Sept. 2022 – Fév. 2023")
org_line("Université Grenoble Alpes — projet d'équipe (3 étudiants)")
bullet("Réseaux de neurones siamois (Python) pour la classification et la détection de nouveauté sur données de spectrométrie de masse, méthode leave-one-type-out.")

role_line("Analyse spatiale des banques coopératives en France", "Sept. 2021 – Fév. 2022")
org_line("Université Grenoble Alpes — projet d'équipe (5 étudiants)")
bullet("Web scraping des données bancaires, données socio-économiques INSEE, application R Shiny pour la visualisation géographique et l'identification des déserts bancaires.")

section_heading("Formation")
edu_line("Master 1 & 2 Mathématiques et Applications — parcours Statistiques et Science des Données", "2021 – 2023")
edu_org("Université Grenoble Alpes")
edu_line("Licence 3 Mathématiques & Master 1 Mathématiques — parcours Analyse Appliquée et Physique Mathématique", "2018 – 2021")
edu_org("Université de Toulon")
edu_line("Licence Mathématiques, Physique et Informatique", "2017 – 2018")
edu_org("Université Cheikh Anta Diop, Dakar")

section_heading("Compétences techniques")
skill_line("Langages", "Python, R, SQL, Julia, JavaScript, HTML/CSS")
skill_line("Web & Frameworks", "Django, Tailwind CSS, HTMX, Bootstrap")
skill_line("Data & Machine Learning", "Machine learning, deep learning (réseaux de neurones), statistiques (tests d'hypothèse, modèles linéaires, séries temporelles, statistique bayésienne)")
skill_line("Bases de données", "PostgreSQL, MySQL, SQLite")
skill_line("Business Intelligence", "Power BI, Qlik Sense, Tableau, R Shiny")
skill_line("DevOps & Outils", "Docker, Git/GitHub/GitLab, intégration continue (GitHub Actions), Render, WeasyPrint, LaTeX")

section_heading("Langues")
plain("Français — langue maternelle   |   Anglais — B2   |   Wolof   |   Espagnol", after=2)

section_heading("Centres d'intérêt")
plain("Sport et activités physiques, lecture, voyages, veille technologique.", after=0)

doc.save("/tmp/cvbuild/cv_cheikh_darou_beye.docx")
print("saved")
